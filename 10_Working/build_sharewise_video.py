"""Build the English ShareWise v9 slide video and its reviewable storyboard.

The source script is public marketing copy. Speech requests send only each
scene's narration text and voice instructions. No ShareWise document is sent.
All rendering and audio caches stay in the system temporary directory.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
import textwrap
import wave

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "10_Working" / "sharewise_video_script_en_original_v9.json"
STORYBOARD = ROOT / "10_Working" / "sharewise_video_storyboard_en_original_v9.pptx"
VIDEO = ROOT / "sharewise" / "sharewise-demo-en-v9.mp4"
POSTER = ROOT / "sharewise" / "sharewise-demo-poster-v9.jpg"
CAPTIONS = ROOT / "sharewise" / "sharewise-demo-en-v9.vtt"
SCRATCH = Path(tempfile.gettempdir()) / "alpinno-sharewise-v9-video"

NAVY = RGBColor(11, 31, 58)
NAVY_2 = RGBColor(21, 53, 86)
BLUE = RGBColor(23, 105, 170)
GOLD = RGBColor(215, 162, 74)
RED = RGBColor(207, 48, 44)
PAPER = RGBColor(247, 248, 247)
WHITE = RGBColor(255, 255, 255)
MIST = RGBColor(235, 242, 247)
LINE = RGBColor(203, 216, 227)
MUTED = RGBColor(79, 100, 119)


def run(args: list[str | Path], *, capture: bool = True) -> subprocess.CompletedProcess:
    return subprocess.run(
        [str(arg) for arg in args], check=True, text=True,
        stdout=subprocess.PIPE if capture else None,
        stderr=subprocess.PIPE if capture else None,
    )


def box(slide, x, y, w, h, fill=WHITE, line=LINE, radius=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background() if line is None else None
    if line is not None:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    return shape


def txt(slide, value, x, y, w, h, size=22, color=NAVY, bold=False, align=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.02)
    frame.margin_top = frame.margin_bottom = Inches(0.02)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for index, line in enumerate(value.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = "Arial"
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color
        if align is not None:
            paragraph.alignment = align
    return shape


def base_slide(prs, number, title, *, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY if dark else PAPER
    box(slide, 0, 0, 13.333, 0.15, GOLD if dark else NAVY, None, False)
    if not dark:
        txt(slide, title, 0.58, 0.47, 12.0, 0.65, 32, NAVY, True)
        txt(slide, "SHAREWISE / ALPINNO GmbH", 0.58, 7.14, 5.0, 0.18, 9, MUTED)
        txt(slide, f"{number:02d}", 12.35, 7.14, 0.4, 0.18, 9, MUTED, align=PP_ALIGN.RIGHT)
    return slide


def label(slide, value, x, y, w, color=BLUE):
    box(slide, x, y, w, 0.34, color, None)
    txt(slide, value, x + 0.08, y + 0.035, w - 0.16, 0.24, 10, WHITE, True, PP_ALIGN.CENTER)


def card(slide, x, y, w, h, heading, body, accent=BLUE):
    box(slide, x, y, w, h)
    box(slide, x, y, 0.065, h, accent, None, False)
    txt(slide, heading, x + 0.25, y + 0.22, w - 0.52, 0.48, 22, NAVY, True)
    txt(slide, body, x + 0.25, y + 1.00, w - 0.52, min(1.60, h - 1.20), 18, NAVY_2)


def add_speaker_notes(slide, scene):
    try:
        slide.notes_slide.notes_text_frame.text = (
            scene["text"] + "\n\nClaim scope: " + scene["scope"]
        )
    except AttributeError:
        pass


def build_storyboard(config):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    scenes = config["scenes"]

    # 1: Product loop
    s = base_slide(prs, 1, scenes[0]["title"], dark=True)
    txt(s, "ShareWise", 0.62, 0.58, 7.5, 0.85, 42, WHITE, True)
    txt(s, "Before AI, prepare privately.", 0.62, 1.44, 10.9, 0.8, 37, WHITE, True)
    txt(s, "A local step between your document and the AI you choose.", 0.63, 2.28, 11.3, 0.56, 22, RGBColor(218, 232, 243))
    for n, x, head, body in [
        (1, 0.62, "Prepare", "Hide detected details locally"),
        (2, 4.86, "Choose", "Review the text you may share"),
        (3, 9.10, "Restore", "Optional, with your private file"),
    ]:
        box(s, x, 3.46, 3.6, 2.14, WHITE, None)
        txt(s, f"0{n}", x + 0.25, 3.70, 0.58, 0.44, 20, GOLD, True)
        txt(s, head, x + 0.25, 4.18, 3.05, 0.48, 24, NAVY, True)
        txt(s, body, x + 0.25, 4.78, 3.04, 0.56, 15, NAVY_2)
    txt(s, "ALPINNO GmbH  |  Developed in Switzerland", 0.62, 6.88, 8.0, 0.24, 13, GOLD, True)
    add_speaker_notes(s, scenes[0])

    # 2: Offline and familiar formats
    s = base_slide(prs, 2, scenes[1]["title"])
    box(s, 0.60, 1.47, 4.33, 4.76)
    txt(s, "98%", 0.90, 1.76, 3.5, 1.32, 70, GOLD, True)
    txt(s, "token recall", 0.90, 3.20, 3.4, 0.50, 24, NAVY, True)
    txt(s, "Text-detector result in a synthetic benchmark. Document results vary.", 0.90, 4.06, 3.46, 1.05, 17, MUTED)
    box(s, 5.23, 1.47, 7.50, 4.76)
    txt(s, "No upload. No account. No model server.", 5.57, 1.76, 6.80, 0.9, 24, NAVY, True)
    txt(s, "Common inputs", 5.57, 3.00, 6.2, 0.40, 18, BLUE, True)
    txt(s, "DOCX  ·  PPTX  ·  XLSX\nText PDF  ·  TXT  ·  Markdown  ·  Pasted text", 5.57, 3.48, 6.45, 1.48, 20, NAVY_2)
    txt(s, "ShareWise preparation runs on your Windows computer.", 0.66, 6.48, 11.9, 0.40, 20, NAVY, True)
    add_speaker_notes(s, scenes[1])

    # 3: OCR support
    s = base_slide(prs, 3, scenes[2]["title"])
    txt(s, "Smart OCR turns printed pages into text for review.", 0.65, 1.36, 11.9, 0.59, 23, NAVY_2)
    card(s, 0.63, 2.22, 3.82, 2.52, "Scanned PDF", "Printed scans and mixed PDFs with native text", GOLD)
    card(s, 4.76, 2.22, 3.82, 2.52, "PNG and JPEG", "Printed page images and photographs", BLUE)
    card(s, 8.89, 2.22, 3.82, 2.52, "Multipage TIFF", "Several printed pages in one image file", RED)
    box(s, 0.63, 5.14, 12.08, 1.22, MIST, None)
    txt(s, "The source image stays intact. ShareWise prepares extracted text.", 0.92, 5.43, 11.45, 0.55, 22, NAVY, True)
    add_speaker_notes(s, scenes[2])

    # 4: Mandatory OCR source confirmation
    s = base_slide(prs, 4, scenes[3]["title"])
    box(s, 0.63, 1.45, 5.93, 3.72)
    label(s, "SOURCE PAGE", 0.91, 1.73, 1.64)
    txt(s, "CONTACT REGISTER", 0.91, 2.30, 5.2, 0.43, 21, NAVY, True)
    txt(s, "Name       Maya Chen\nEmail       maya.chen@example.com\nAccount   ACCT-EXAMPLE-4829", 0.91, 2.96, 5.18, 1.48, 18, NAVY_2)
    box(s, 6.78, 1.45, 5.93, 3.72)
    label(s, "RECOGNIZED TEXT", 7.06, 1.73, 2.35, GOLD)
    txt(s, "Compare and correct", 7.06, 2.30, 5.16, 0.48, 21, NAVY, True)
    txt(s, "Maya Chen\nmaya.chen@example.com\nACCT-EXAMPLE-4829", 7.06, 2.96, 5.10, 1.43, 18, NAVY_2)
    txt(s, "Every OCR page needs your confirmation before privacy analysis.", 0.67, 5.43, 11.97, 0.50, 20, NAVY, True)
    txt(s, "OCR document languages: English · Italian · French · German · Spanish · Dutch", 0.67, 6.10, 11.98, 0.39, 15, BLUE, True)
    add_speaker_notes(s, scenes[3])

    # 5: Large paged documents
    s = base_slide(prs, 5, scenes[4]["title"])
    for x, figure, unit, sub in [
        (0.63, "200", "pages", "per paged document"),
        (4.76, "512", "MiB", "input size"),
        (8.89, "2 million", "characters", "recognized or extracted"),
    ]:
        box(s, x, 1.69, 3.82, 3.29)
        txt(s, figure, x + 0.22, 2.00, 3.38, 0.90, 42, GOLD, True, PP_ALIGN.CENTER)
        txt(s, unit, x + 0.22, 3.01, 3.38, 0.46, 23, NAVY, True, PP_ALIGN.CENTER)
        txt(s, sub, x + 0.22, 3.70, 3.38, 0.60, 15, MUTED, align=PP_ALIGN.CENTER)
    txt(s, "Applies to paged PDFs and images. Other formats have separate limits.", 0.67, 5.31, 11.98, 0.50, 20, NAVY, True)
    txt(s, "Incomplete or failed work is reported, not silently approved for export.", 0.67, 5.97, 11.98, 0.48, 18, BLUE)
    add_speaker_notes(s, scenes[4])

    # 6: Modes and human review
    s = base_slide(prs, 6, scenes[5]["title"])
    card(s, 0.63, 1.50, 5.89, 3.76, "Minimal", "A quick path for standard inputs. Scanned pages pause for the required OCR source check.", BLUE)
    card(s, 6.81, 1.50, 5.89, 3.76, "Advanced", "Aligned text comparison, batch work and controls to keep or hide suggestions.", GOLD)
    box(s, 0.63, 5.64, 12.07, 0.86, MIST, None)
    txt(s, "Review the complete output, including text that was not highlighted.", 0.92, 5.82, 11.55, 0.48, 20, NAVY, True)
    add_speaker_notes(s, scenes[5])

    # 7: Exact lists and local memory
    s = base_slide(prs, 7, scenes[6]["title"])
    card(s, 0.63, 1.68, 5.89, 3.78, "Confidential-term lists", "Add exact client, company or project terms you want hidden in extracted text.", BLUE)
    card(s, 6.81, 1.68, 5.89, 3.78, "Correction memory", "Remember your deliberate keep and hide choices locally. Disable or clear them in Settings.", GOLD)
    txt(s, "No cloud learning. No inference from untouched suggestions.", 0.66, 5.87, 11.9, 0.49, 20, NAVY, True)
    add_speaker_notes(s, scenes[6])

    # 8: Restoration
    s = base_slide(prs, 8, scenes[7]["title"])
    box(s, 0.63, 1.58, 4.31, 3.40)
    label(s, "AI RESPONSE", 0.89, 1.84, 1.79)
    txt(s, "Dear [PERSON_1],\nthank you for the update.", 0.89, 2.53, 3.73, 1.28, 21, NAVY_2)
    box(s, 5.14, 2.89, 1.15, 0.38, GOLD, None)
    txt(s, "→", 5.45, 2.87, 0.55, 0.40, 23, NAVY, True, PP_ALIGN.CENTER)
    box(s, 6.47, 1.58, 6.23, 3.40)
    label(s, "RESTORED LOCALLY", 6.75, 1.84, 2.30, BLUE)
    txt(s, "Dear Maya Chen,\nthank you for the update.", 6.75, 2.53, 5.53, 1.28, 21, NAVY_2)
    txt(s, "Keep the matching restoration file private. It holds original values and is not encrypted.", 0.67, 5.39, 11.94, 0.97, 20, NAVY, True)
    add_speaker_notes(s, scenes[7])

    # 9: Final responsibility and contact
    s = base_slide(prs, 9, scenes[8]["title"], dark=True)
    txt(s, "ShareWise", 0.63, 0.65, 8.1, 0.75, 38, WHITE, True)
    txt(s, "Your documents.\nYour rules.\nYour decision.", 0.63, 1.62, 8.0, 3.20, 48, WHITE, True)
    box(s, 8.87, 1.92, 3.81, 2.96, WHITE, None)
    txt(s, "Review before sharing", 9.18, 2.25, 3.23, 0.94, 24, NAVY, True)
    txt(s, "Detection helps. It cannot guarantee every private detail was found.", 9.18, 3.30, 3.23, 1.22, 18, NAVY_2)
    txt(s, "Ask ALPINNO about availability and licensing.", 0.66, 5.66, 11.93, 0.62, 23, GOLD, True)
    txt(s, "ALPINNO GmbH  |  Developed in Switzerland", 0.66, 6.92, 8.5, 0.24, 13, RGBColor(218, 232, 243))
    add_speaker_notes(s, scenes[8])

    STORYBOARD.parent.mkdir(parents=True, exist_ok=True)
    prs.save(STORYBOARD)
    return len(prs.slides)


def render_slides(count):
    SCRATCH.mkdir(parents=True, exist_ok=True)
    pdf = SCRATCH / STORYBOARD.with_suffix(".pdf").name
    soffice = shutil.which("soffice")
    if not soffice:
        raise RuntimeError("LibreOffice is unavailable")
    soffice_command = ["cmd.exe", "/c", soffice] if Path(soffice).suffix.lower() in {".cmd", ".bat"} else [soffice]
    profile = (SCRATCH / "libreoffice-profile").resolve().as_uri()
    run([*soffice_command, f"-env:UserInstallation={profile}", "--headless", "--convert-to", "pdf", "--outdir", SCRATCH, STORYBOARD])
    run(["pdftoppm", "-png", "-r", "144", pdf, SCRATCH / "slide"])
    slides = [SCRATCH / f"slide-{n}.png" for n in range(1, count + 1)]
    if not all(path.is_file() for path in slides):
        raise RuntimeError("Storyboard rendering did not produce every slide")
    for path in slides:
        with Image.open(path) as image:
            if image.size != (1920, 1080):
                raise RuntimeError(f"Unexpected rendered slide size: {path.name} {image.size}")
    with Image.open(slides[0]) as image:
        image.convert("RGB").save(POSTER, format="JPEG", quality=88, optimize=True)
    return slides


def speech(config, scenes):
    if not os.environ.get("OPENAI_API_KEY"):
        raise RuntimeError("OPENAI_API_KEY is unavailable")
    from openai import OpenAI

    client = OpenAI(api_key=os.environ["OPENAI_API_KEY"], max_retries=0, timeout=90)
    tracks = []
    try:
        for index, scene in enumerate(scenes, 1):
            payload = {
                "model": config["voice_model"],
                "voice": config["voice"],
                "instructions": config["voice_instructions"],
                "speed": 1,
                "input": scene["text"],
                "response_format": "wav",
            }
            fingerprint = hashlib.sha256(json.dumps(payload, sort_keys=True).encode()).hexdigest()[:16]
            target = SCRATCH / f"voice-{index:02d}-{fingerprint}.wav"
            if not target.exists():
                partial = target.with_suffix(".partial.wav")
                with client.audio.speech.with_streaming_response.create(**payload) as response:
                    response.stream_to_file(partial)
                if partial.stat().st_size < 10000:
                    raise RuntimeError(f"Narration segment {index} was unexpectedly short")
                partial.replace(target)
            tracks.append(target)
            print(f"Narration segment {index}/{len(scenes)} ready", flush=True)
    finally:
        client.close()
    return tracks


def join_audio(tracks):
    joined = SCRATCH / "joined.wav"
    timings = []
    lead, gap, tail = 0.6, 0.48, 1.6
    cursor = lead
    with wave.open(str(joined), "wb") as out:
        format_ = None
        for index, source in enumerate(tracks):
            with wave.open(str(source), "rb") as inp:
                spec = (inp.getnchannels(), inp.getsampwidth(), inp.getframerate())
                if format_ is None:
                    format_ = spec
                    out.setnchannels(spec[0])
                    out.setsampwidth(spec[1])
                    out.setframerate(spec[2])
                    out.writeframes(b"\0" * round(lead * spec[2]) * spec[0] * spec[1])
                if spec != format_:
                    raise RuntimeError("Speech segments use inconsistent WAV formats")
                # Streamed WAV headers may advertise a placeholder frame count.
                # Use the bytes actually returned by the file instead.
                frames = inp.readframes(inp.getnframes())
                duration = len(frames) / (spec[0] * spec[1] * spec[2])
                out.writeframes(frames)
                timings.append((cursor, cursor + duration))
                cursor += duration
                pause = tail if index == len(tracks) - 1 else gap
                out.writeframes(b"\0" * round(pause * spec[2]) * spec[0] * spec[1])
                cursor += pause
    return joined, timings, cursor


def stamp(seconds):
    milliseconds = round(seconds * 1000)
    hours, remaining = divmod(milliseconds, 3600000)
    minutes, remaining = divmod(remaining, 60000)
    secs, milliseconds = divmod(remaining, 1000)
    return f"{hours:02d}:{minutes:02d}:{secs:02d}.{milliseconds:03d}"


def write_captions(scenes, timings):
    cues = ["WEBVTT", ""]
    for scene, (start, end) in zip(scenes, timings):
        sentences = re.split(r"(?<=[.!?])\s+", scene["text"])
        weights = [max(1, len(sentence.split())) for sentence in sentences]
        total = sum(weights)
        cursor = start
        for index, (sentence, weight) in enumerate(zip(sentences, weights)):
            until = end if index == len(sentences) - 1 else cursor + (end - start) * weight / total
            cues.extend([f"{stamp(cursor)} --> {stamp(until)}", textwrap.fill(sentence, width=52), ""])
            cursor = until
    CAPTIONS.write_text("\n".join(cues), encoding="utf-8")


def encode(slides, joined, timings, duration):
    boundaries = [0.0]
    boundaries += [(timings[i - 1][1] + timings[i][0]) / 2 for i in range(1, len(timings))]
    boundaries.append(duration)
    lines = ["ffconcat version 1.0"]
    for index, path in enumerate(slides):
        safe = path.resolve().as_posix()
        if "'" in safe:
            raise RuntimeError("Apostrophe in slideshow path is unsupported")
        lines.extend([f"file '{safe}'", f"duration {boundaries[index + 1] - boundaries[index]:.6f}"])
    lines.append(f"file '{slides[-1].resolve().as_posix()}'")
    concat = SCRATCH / "slides.ffconcat"
    concat.write_text("\n".join(lines) + "\n", encoding="utf-8")
    run([
        "ffmpeg", "-hide_banner", "-loglevel", "error", "-y",
        "-f", "concat", "-safe", "0", "-i", concat,
        "-i", joined, "-map", "0:v:0", "-map", "1:a:0",
        "-vf", "fps=30,format=yuv420p", "-c:v", "libx264", "-preset", "medium",
        "-tune", "stillimage", "-crf", "20", "-c:a", "aac", "-b:a", "160k",
        "-af", "loudnorm=I=-16:TP=-1.5:LRA=9", "-t", f"{duration:.4f}",
        "-movflags", "+faststart", "-metadata", "title=ShareWise English demonstration v9",
        "-metadata", "comment=AI-generated narration; fictional illustrative visuals", VIDEO,
    ])
    run(["ffmpeg", "-v", "error", "-i", VIDEO, "-f", "null", "-"])
    probe = json.loads(run(["ffprobe", "-v", "error", "-show_format", "-show_streams", "-of", "json", VIDEO]).stdout)
    video = next(stream for stream in probe["streams"] if stream["codec_type"] == "video")
    audio = next(stream for stream in probe["streams"] if stream["codec_type"] == "audio")
    if (video["width"], video["height"], video["codec_name"], audio["codec_name"]) != (1920, 1080, "h264", "aac"):
        raise RuntimeError("Unexpected video codec or dimensions")
    return float(probe["format"]["duration"])


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--render-only", action="store_true", help="Build and render slides without API use")
    args = parser.parse_args()
    config = json.loads(SCRIPT.read_text(encoding="utf-8"))
    scenes = config["scenes"]
    if len(scenes) != 9 or any("\u2014" in scene["text"] for scene in scenes):
        raise RuntimeError("Unexpected script shape or punctuation")
    count = build_storyboard(config)
    slides = render_slides(count)
    if args.render_only:
        print(f"Rendered {count} storyboard slides in {SCRATCH}")
        return
    tracks = speech(config, scenes)
    joined, timings, duration = join_audio(tracks)
    write_captions(scenes, timings)
    actual = encode(slides, joined, timings, duration)
    print(json.dumps({"video": str(VIDEO), "seconds": actual, "poster": str(POSTER), "captions": str(CAPTIONS)}))


if __name__ == "__main__":
    main()
